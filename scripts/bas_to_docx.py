#!/usr/bin/env python3
"""Build a .docx directly with python-docx: centered step diagram + monospace code.

Going straight to .docx avoids HTML-import quirks entirely. Key tricks:
  * the diagram is a 1-column table centered with WD_TABLE_ALIGNMENT.CENTER,
  * paragraph spacing is set explicitly (space_before/after = 0, exact line spacing)
    so code lines don't get Word's default "space after",
  * code uses Courier New and cell shading is applied via w:shd.

Usage:
  python scripts/bas_to_docx.py samples/grade_calculator.bas \
      --output-text samples/grade_calculator_output.txt -o output/grade_calculator.docx
"""
from __future__ import annotations

import argparse
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from qb64_parse import parse_steps, program_title  # noqa: E402

BLUE = RGBColor(0x1F, 0x38, 0x64)


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def borders(cell, color: str = "2F5496", size: int = 12) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))  # eighths of a point
        b.set(qn("w:color"), color)
        el.append(b)
    tc_pr.append(el)


def tight(paragraph, line_pt: float | None = None, after: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    if line_pt:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(line_pt)


def add_diagram(doc, steps) -> None:
    rows = len(steps) * 2 - 1
    table = doc.add_table(rows=rows, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, step in enumerate(steps):
        cell = table.cell(i * 2, 0)
        cell.width = Pt(240)
        shade(cell, "DEEAF6")
        borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 14)
        p.add_run(f"Step {step.number}").bold = True
        p2 = cell.add_paragraph(step.title)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p2, 14)
        if i < len(steps) - 1:
            arrow_cell = table.cell(i * 2 + 1, 0)
            arrow_cell.width = Pt(240)
            ap = arrow_cell.paragraphs[0]
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tight(ap, 16)
            run = ap.add_run("\u25BC")
            run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
    doc.add_paragraph()


def add_code(doc, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F4F4F4")
    borders(cell, "CCCCCC", 4)
    lines = code.splitlines() or [""]
    for n, line in enumerate(lines):
        p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        tight(p, 12)
        run = p.add_run(line)  # spaces are preserved in .docx runs
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(10)
    tight(doc.add_paragraph(), after=6)


def build_docx(source: str, output_text: str | None, out_path: str) -> int:
    steps = parse_steps(source)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    h = doc.add_heading(program_title(source), level=1)
    h.runs[0].font.color.rgb = BLUE
    doc.add_paragraph("A step-by-step walkthrough of a QB64 program.")
    doc.add_heading("Program flow", level=2)
    add_diagram(doc, steps)
    doc.add_heading("Step-by-step code", level=2)
    for step in steps:
        doc.add_heading(f"Step {step.number}: {step.title}", level=3)
        add_code(doc, step.code)
    if output_text is not None:
        doc.add_heading("Program output", level=2)
        add_code(doc, output_text.rstrip("\n"))
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return len(steps)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("source")
    ap.add_argument("--output-text")
    ap.add_argument("-o", "--out", default="output/program.docx")
    args = ap.parse_args(argv)
    with open(args.source, encoding="utf-8") as fh:
        source = fh.read()
    output_text = None
    if args.output_text:
        with open(args.output_text, encoding="utf-8") as fh:
            output_text = fh.read()
    n = build_docx(source, output_text, args.out)
    print(f"wrote {args.out} ({n} steps)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
