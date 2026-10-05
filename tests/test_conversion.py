import os
import re
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

from bas_to_html import build_html, highlight, preserve_spaces  # noqa: E402
from qb64_parse import parse_steps, program_title  # noqa: E402

SAMPLE = os.path.join(ROOT, "samples", "grade_calculator.bas")
OUTPUT = os.path.join(ROOT, "samples", "grade_calculator_output.txt")


def read(path):
    with open(path, encoding="utf-8") as fh:
        return fh.read()


class ParserTests(unittest.TestCase):
    def test_steps_found_in_order(self):
        steps = parse_steps(read(SAMPLE))
        self.assertEqual([s.number for s in steps], [1, 2, 3, 4, 5])
        self.assertEqual(steps[0].title, "Declare variables")
        self.assertTrue(steps[1].code.startswith("DATA 78,92,85,64,88"))

    def test_title(self):
        self.assertEqual(program_title(read(SAMPLE)), "Grade Calculator")

    def test_sample_output_matches_program_logic(self):
        scores = [78, 92, 85, 64, 88]
        out = read(OUTPUT)
        self.assertIn(f"Total:    {sum(scores)}", out)
        self.assertIn(f"Average:  {sum(scores) / 5:.2f}", out)
        self.assertIn("Letter grade: B", out)


class HtmlTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = build_html(read(SAMPLE), read(OUTPUT))

    def test_word_safe_css_only(self):
        for bad in ("display:flex", "display:grid", "<style", "class=", "<link", "var(--"):
            self.assertNotIn(bad, self.doc)

    def test_centered_table_diagram(self):
        self.assertIn('<table align="center"', self.doc)
        self.assertEqual(len(re.findall(r"<b>Step \d</b>", self.doc)), 5)
        self.assertEqual(self.doc.count("&#9660;"), 4)  # arrows between 5 boxes

    def test_code_is_monospace_and_tight(self):
        self.assertIn("'Courier New'", self.doc)
        self.assertIn("margin:0;line-height:12pt;mso-line-height-rule:exactly", self.doc)

    def test_escaping_and_highlight(self):
        self.assertIn("&gt;=", self.doc)  # CASE IS >= 90
        self.assertIn('<span style="color:#A31515;">&quot;A&quot;</span>', highlight('grade = "A"'))

    def test_spaces_preserved_outside_tags_only(self):
        frag = preserve_spaces('    <span style="color:red;">X  Y</span>')
        self.assertEqual(frag, '&nbsp;&nbsp;&nbsp;&nbsp;<span style="color:red;">X&nbsp;&nbsp;Y</span>')


class DocxTests(unittest.TestCase):
    def test_docx_has_centered_diagram(self):
        try:
            import docx
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from bas_to_docx import build_docx
        except ImportError:
            self.skipTest("python-docx not installed")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.docx")
            self.assertEqual(build_docx(read(SAMPLE), read(OUTPUT), path), 5)
            d = docx.Document(path)
            diagram = d.tables[0]
            self.assertEqual(diagram.alignment, WD_TABLE_ALIGNMENT.CENTER)
            self.assertEqual(len(diagram.rows), 9)
            self.assertEqual(diagram.cell(0, 0).paragraphs[0].text, "Step 1")
            code_font = d.tables[1].cell(0, 0).paragraphs[0].runs[0].font.name
            self.assertEqual(code_font, "Courier New")


if __name__ == "__main__":
    unittest.main()
